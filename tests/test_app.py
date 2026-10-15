"""Integration tests for Question Paper Alignment Software.

Tests API endpoints, WeasyPrint PDF pipeline, python-docx export pipeline,
security restrictions, styling, figures, and schema validation.
"""

import base64
import io
import json
import subprocess
import sys
import time
import unittest
import urllib.error
import urllib.request
from pathlib import Path
from docx import Document

BASE = Path(__file__).resolve().parent.parent

# Small 1x1 valid PNG for testing
TINY_PNG_B64 = (
    'data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=='
)


class IntegrationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.port = 8123
        cls.base_url = f'http://127.0.0.1:{cls.port}'
        cls.proc = subprocess.Popen(
            [sys.executable, '-m', 'uvicorn', 'app.main:app', '--port', str(cls.port)],
            cwd=str(BASE),
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
        )
        max_retries = 25
        for _ in range(max_retries):
            try:
                with urllib.request.urlopen(f'{cls.base_url}/health', timeout=1) as resp:
                    if resp.status == 200:
                        break
            except Exception:
                time.sleep(0.3)
        else:
            cls.proc.terminate()
            raise RuntimeError('Test server failed to start within timeout.')

    @classmethod
    def tearDownClass(cls):
        if cls.proc:
            cls.proc.terminate()
            cls.proc.wait()

    def test_health_check(self):
        with urllib.request.urlopen(f'{self.base_url}/health') as resp:
            self.assertEqual(resp.status, 200)
            data = json.loads(resp.read().decode('utf-8'))
            self.assertEqual(data.get('status'), 'ok')

    def test_options_endpoint(self):
        with urllib.request.urlopen(f'{self.base_url}/options') as resp:
            self.assertEqual(resp.status, 200)
            data = json.loads(resp.read().decode('utf-8'))
            self.assertIn('fonts', data)
            self.assertIn('Times New Roman', data['fonts'])
            self.assertIn('Arial', data['fonts'])
            self.assertIn('sizes', data)
            self.assertIn(11, data['sizes'])
            self.assertIn('formats', data)
            self.assertEqual(data['formats'], ['pdf', 'docx'])
            self.assertIn('figure', data)
            self.assertEqual(data['figure']['maxBytes'], 1572864)
            self.assertEqual(data['figure']['maxCount'], 12)
            self.assertIn(50, data['figure']['widths'])

    def test_frontend_index_served(self):
        with urllib.request.urlopen(f'{self.base_url}/') as resp:
            self.assertEqual(resp.status, 200)
            html = resp.read().decode('utf-8')
            self.assertIn('Question Paper Alignment Software', html)
            self.assertIn('pane-editor', html)
            self.assertIn('pane-preview', html)

    def test_frontend_static_assets(self):
        with urllib.request.urlopen(f'{self.base_url}/css/tokens.css') as resp:
            self.assertEqual(resp.status, 200)
            css = resp.read().decode('utf-8')
            self.assertIn(':root', css)

        with urllib.request.urlopen(f'{self.base_url}/js/app.js') as resp:
            self.assertEqual(resp.status, 200)
            js = resp.read().decode('utf-8')
            self.assertIn('ToolbarComponent', js)

    def test_preview_endpoint(self):
        sample_path = BASE / 'tests' / 'sample_paper.json'
        payload = sample_path.read_bytes()
        req = urllib.request.Request(
            f'{self.base_url}/preview',
            data=payload,
            headers={'Content-Type': 'application/json'},
            method='POST',
        )
        with urllib.request.urlopen(req) as resp:
            self.assertEqual(resp.status, 200)
            html = resp.read().decode('utf-8')
            self.assertIn('ST JOHN THE BAPTIST', html)
            self.assertIn('Section A', html)
            self.assertIn('--paper-font', html)
            self.assertIn('--paper-size', html)

    def test_generate_pdf_endpoint(self):
        sample_path = BASE / 'tests' / 'sample_paper.json'
        payload = sample_path.read_bytes()
        req = urllib.request.Request(
            f'{self.base_url}/generate?format=pdf',
            data=payload,
            headers={'Content-Type': 'application/json'},
            method='POST',
        )
        with urllib.request.urlopen(req) as resp:
            self.assertEqual(resp.status, 200)
            self.assertEqual(resp.headers.get_content_type(), 'application/pdf')
            self.assertIn('.pdf', resp.headers.get('Content-Disposition', ''))
            pdf_data = resp.read()
            self.assertTrue(pdf_data.startswith(b'%PDF'))

    def test_generate_docx_endpoint(self):
        sample_path = BASE / 'tests' / 'sample_paper.json'
        payload = sample_path.read_bytes()
        req = urllib.request.Request(
            f'{self.base_url}/generate?format=docx',
            data=payload,
            headers={'Content-Type': 'application/json'},
            method='POST',
        )
        with urllib.request.urlopen(req) as resp:
            self.assertEqual(resp.status, 200)
            self.assertEqual(
                resp.headers.get_content_type(),
                'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            )
            self.assertIn('.docx', resp.headers.get('Content-Disposition', ''))
            docx_data = resp.read()
            self.assertTrue(docx_data.startswith(b'PK'))

            # Acceptance Criterion 4: Verify reopening docx with python-docx
            doc = Document(io.BytesIO(docx_data))
            self.assertTrue(len(doc.paragraphs) > 0)
            self.assertTrue(len(doc.tables) > 0)

            # Check continuous question numbering matches
            all_text = ' '.join(p.text for p in doc.paragraphs)
            self.assertIn('1.', all_text)
            self.assertIn('Section A', all_text)

    def test_dev_sample_endpoints(self):
        with urllib.request.urlopen(f'{self.base_url}/dev/sample') as resp:
            self.assertEqual(resp.status, 200)
            html = resp.read().decode('utf-8')
            self.assertIn('ST JOHN THE BAPTIST', html)

        with urllib.request.urlopen(f'{self.base_url}/dev/sample.pdf') as resp:
            self.assertEqual(resp.status, 200)
            pdf_data = resp.read()
            self.assertTrue(pdf_data.startswith(b'%PDF'))

    def test_style_changes_preview_pdf_and_docx(self):
        # Acceptance Criterion 2: Changing font and size changes preview, PDF, and DOCX consistently
        paper = {
            'school': 'ACME INTERNATIONAL SCHOOL',
            'exam': 'MIDTERM EXAMINATION',
            'class': 'X',
            'subject': 'PHYSICS',
            'time': '2 hrs',
            'totalMarks': 50,
            'style': {'fontFamily': 'Georgia', 'fontSize': 12},
            'sections': [
                {
                    'label': 'A',
                    'marksPerQuestion': 2,
                    'questions': [{'type': 'short', 'text': 'State Ohm’s law.'}],
                }
            ],
        }
        payload = json.dumps(paper).encode('utf-8')

        # 1. Preview
        req = urllib.request.Request(
            f'{self.base_url}/preview',
            data=payload,
            headers={'Content-Type': 'application/json'},
            method='POST',
        )
        with urllib.request.urlopen(req) as resp:
            self.assertEqual(resp.status, 200)
            html = resp.read().decode('utf-8')
            self.assertIn('Georgia', html)
            self.assertIn('12pt', html)

        # 2. PDF
        req = urllib.request.Request(
            f'{self.base_url}/generate?format=pdf',
            data=payload,
            headers={'Content-Type': 'application/json'},
            method='POST',
        )
        with urllib.request.urlopen(req) as resp:
            self.assertEqual(resp.status, 200)
            pdf = resp.read()
            self.assertTrue(pdf.startswith(b'%PDF'))

        # 3. DOCX
        req = urllib.request.Request(
            f'{self.base_url}/generate?format=docx',
            data=payload,
            headers={'Content-Type': 'application/json'},
            method='POST',
        )
        with urllib.request.urlopen(req) as resp:
            self.assertEqual(resp.status, 200)
            docx_bytes = resp.read()
            doc = Document(io.BytesIO(docx_bytes))
            normal_font = doc.styles['Normal'].font
            self.assertEqual(normal_font.name, 'Georgia')
            self.assertEqual(normal_font.size.pt, 12.0)

    def test_paper_with_png_figure_on_short_question(self):
        # Acceptance Criterion 3: A paper with a PNG figure on a short question shows the figure in preview, PDF, and DOCX
        paper = {
            'school': 'MODEL ACADEMY',
            'exam': 'FINAL EXAM',
            'class': 'VIII',
            'subject': 'MATHEMATICS',
            'time': '3 hrs',
            'totalMarks': 80,
            'sections': [
                {
                    'label': 'A',
                    'marksPerQuestion': 5,
                    'questions': [
                        {
                            'type': 'short',
                            'text': 'Find the area of the shaded triangle in the figure below.',
                            'figure': {
                                'src': TINY_PNG_B64,
                                'widthPct': 65,
                                'align': 'center',
                                'caption': 'Fig. 1 Triangle ABC',
                            },
                        }
                    ],
                }
            ],
        }
        payload = json.dumps(paper).encode('utf-8')

        # 1. Preview
        req = urllib.request.Request(
            f'{self.base_url}/preview',
            data=payload,
            headers={'Content-Type': 'application/json'},
            method='POST',
        )
        with urllib.request.urlopen(req) as resp:
            self.assertEqual(resp.status, 200)
            html = resp.read().decode('utf-8')
            self.assertIn('q-figure-box', html)
            self.assertIn('Fig. 1 Triangle ABC', html)
            self.assertIn('width: 65%', html)

        # 2. PDF
        req = urllib.request.Request(
            f'{self.base_url}/generate?format=pdf',
            data=payload,
            headers={'Content-Type': 'application/json'},
            method='POST',
        )
        with urllib.request.urlopen(req) as resp:
            self.assertEqual(resp.status, 200)
            pdf = resp.read()
            self.assertTrue(pdf.startswith(b'%PDF'))

        # 3. DOCX
        req = urllib.request.Request(
            f'{self.base_url}/generate?format=docx',
            data=payload,
            headers={'Content-Type': 'application/json'},
            method='POST',
        )
        with urllib.request.urlopen(req) as resp:
            self.assertEqual(resp.status, 200)
            docx_bytes = resp.read()
            doc = Document(io.BytesIO(docx_bytes))
            all_text = ' '.join(p.text for p in doc.paragraphs)
            self.assertIn('Fig. 1 Triangle ABC', all_text)

    def test_invalid_figures_and_styles_return_422(self):
        # Acceptance Criterion 5: Invalid figures and styles return 422 with readable errors
        # 1. Invalid style
        paper = {'style': {'fontFamily': 'Comic Sans', 'fontSize': 72}, 'sections': []}
        req = urllib.request.Request(
            f'{self.base_url}/preview',
            data=json.dumps(paper).encode('utf-8'),
            headers={'Content-Type': 'application/json'},
            method='POST',
        )
        with self.assertRaises(urllib.error.HTTPError) as ctx:
            urllib.request.urlopen(req)
        self.assertEqual(ctx.exception.code, 422)
        err = json.loads(ctx.exception.read().decode('utf-8'))
        self.assertIn('detail', err)
        locs = [d['loc'] for d in err['detail']]
        self.assertIn(['style', 'fontFamily'], locs)
        self.assertIn(['style', 'fontSize'], locs)

        # 2. Invalid remote URL in figure (Criterion 7)
        paper2 = {
            'sections': [
                {
                    'label': 'A',
                    'questions': [
                        {
                            'type': 'short',
                            'text': 'Q1',
                            'figure': {'src': 'http://evil.com/leak.png'},
                        }
                    ],
                }
            ]
        }
        req2 = urllib.request.Request(
            f'{self.base_url}/preview',
            data=json.dumps(paper2).encode('utf-8'),
            headers={'Content-Type': 'application/json'},
            method='POST',
        )
        with self.assertRaises(urllib.error.HTTPError) as ctx:
            urllib.request.urlopen(req2)
        self.assertEqual(ctx.exception.code, 422)
        err2 = json.loads(ctx.exception.read().decode('utf-8'))
        self.assertIn('detail', err2)
        self.assertEqual(err2['detail'][0]['loc'], ['sections', 0, 'questions', 0, 'figure', 'src'])

        # 3. SVG data URI rejected
        paper3 = {
            'sections': [
                {
                    'label': 'A',
                    'questions': [
                        {
                            'type': 'short',
                            'text': 'Q1',
                            'figure': {'src': 'data:image/svg+xml;base64,PHN2Zz48L3N2Zz4='},
                        }
                    ],
                }
            ]
        }
        req3 = urllib.request.Request(
            f'{self.base_url}/preview',
            data=json.dumps(paper3).encode('utf-8'),
            headers={'Content-Type': 'application/json'},
            method='POST',
        )
        with self.assertRaises(urllib.error.HTTPError) as ctx:
            urllib.request.urlopen(req3)
        self.assertEqual(ctx.exception.code, 422)

        # 4. Unknown export format rejected with 422
        req4 = urllib.request.Request(
            f'{self.base_url}/generate?format=xlsx',
            data=json.dumps({'sections': []}).encode('utf-8'),
            headers={'Content-Type': 'application/json'},
            method='POST',
        )
        with self.assertRaises(urllib.error.HTTPError) as ctx:
            urllib.request.urlopen(req4)
        self.assertEqual(ctx.exception.code, 422)

    def test_caption_and_question_text_escaping(self):
        # Acceptance Criterion 6: Caption or question text containing quotes, angle brackets or ampersands is escaped everywhere
        tricky_text = 'Prove that <tag> & "quotes" \'single\' are escaped.'
        tricky_caption = 'Fig. <X> & "Y" \'Z\''
        paper = {
            'school': 'ESCAPING & <TEST> SCHOOL',
            'exam': 'EXAM "SPECIAL" & MORE',
            'class': 'X',
            'subject': 'TEST',
            'time': '1 hr',
            'totalMarks': 10,
            'sections': [
                {
                    'label': 'A',
                    'marksPerQuestion': 10,
                    'questions': [
                        {
                            'type': 'short',
                            'text': tricky_text,
                            'figure': {
                                'src': TINY_PNG_B64,
                                'caption': tricky_caption,
                            },
                        }
                    ],
                }
            ],
        }
        payload = json.dumps(paper).encode('utf-8')

        # 1. Preview HTML
        req = urllib.request.Request(
            f'{self.base_url}/preview',
            data=payload,
            headers={'Content-Type': 'application/json'},
            method='POST',
        )
        with urllib.request.urlopen(req) as resp:
            self.assertEqual(resp.status, 200)
            html = resp.read().decode('utf-8')
            self.assertIn('&lt;tag&gt;', html)
            self.assertIn('&amp;', html)
            self.assertNotIn('<tag>', html)
            self.assertIn('&lt;X&gt;', html)

        # 2. PDF
        req = urllib.request.Request(
            f'{self.base_url}/generate?format=pdf',
            data=payload,
            headers={'Content-Type': 'application/json'},
            method='POST',
        )
        with urllib.request.urlopen(req) as resp:
            self.assertEqual(resp.status, 200)
            self.assertTrue(resp.read().startswith(b'%PDF'))

        # 3. DOCX
        req = urllib.request.Request(
            f'{self.base_url}/generate?format=docx',
            data=payload,
            headers={'Content-Type': 'application/json'},
            method='POST',
        )
        with urllib.request.urlopen(req) as resp:
            self.assertEqual(resp.status, 200)
            doc = Document(io.BytesIO(resp.read()))
            all_text = ' '.join(p.text for p in doc.paragraphs)
            self.assertIn('Prove that <tag> & "quotes"', all_text)
            self.assertIn('Fig. <X> & "Y"', all_text)


if __name__ == '__main__':
    unittest.main()
