import base64
import io
import math
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Mm, Pt
from PIL import Image

from .render import build_view


def _make_borderless(table):
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        '<w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/>'
        '<w:insideH w:val="none"/><w:insideV w:val="none"/>'
        '</w:tblBorders>'
    )
    tblPr.append(tblBorders)


def _set_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))


def _setup_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.text = 'Page '
    p.runs[0].font.size = Pt(9)

    # PAGE field
    r1 = p.add_run()
    r1.font.size = Pt(9)
    fld1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    instr1 = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>')
    fld2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>')
    fld3 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    r1._r.extend([fld1, instr1, fld2, fld3])

    r_of = p.add_run(' of ')
    r_of.font.size = Pt(9)

    # NUMPAGES field
    r2 = p.add_run()
    r2.font.size = Pt(9)
    fld4 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    instr2 = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> NUMPAGES </w:instrText>')
    fld5 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>')
    fld6 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    r2._r.extend([fld4, instr2, fld5, fld6])


def _add_figure(doc, fig, keep_with_next=True):
    if not fig or not fig.get('src'):
        return
    src = fig['src']
    if src.startswith('data:image/png;base64,'):
        b64 = src[len('data:image/png;base64,'):]
    elif src.startswith('data:image/jpeg;base64,'):
        b64 = src[len('data:image/jpeg;base64,'):]
    else:
        return

    try:
        img_bytes = base64.b64decode(b64)
    except Exception:
        return

    width_pct = fig.get('widthPct', 50)
    target_width = Mm(180 * width_pct / 100)
    max_height = Mm(90)
    target_height = None

    try:
        with Image.open(io.BytesIO(img_bytes)) as pil_img:
            orig_w, orig_h = pil_img.size
        aspect = orig_h / orig_w if orig_w > 0 else 1.0
        calc_h = target_width * aspect
        if calc_h > max_height:
            target_height = max_height
            target_width = max_height / aspect if aspect > 0 else target_width
        else:
            target_height = calc_h
    except Exception:
        pass

    p_fig = doc.add_paragraph()
    p_fig.paragraph_format.keep_with_next = keep_with_next
    p_fig.paragraph_format.space_before = Pt(4)
    p_fig.paragraph_format.space_after = Pt(2)
    align = fig.get('align', 'center')
    p_fig.alignment = WD_ALIGN_PARAGRAPH.CENTER if align == 'center' else WD_ALIGN_PARAGRAPH.LEFT
    run_img = p_fig.add_run()
    if target_height:
        run_img.add_picture(io.BytesIO(img_bytes), width=target_width, height=target_height)
    else:
        run_img.add_picture(io.BytesIO(img_bytes), width=target_width)

    caption = fig.get('caption')
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.keep_with_next = keep_with_next
        p_cap.paragraph_format.space_before = Pt(0)
        p_cap.paragraph_format.space_after = Pt(4)
        r_cap = p_cap.add_run(caption)
        r_cap.italic = True
        r_cap.font.size = Pt(9)


def render_docx(paper: dict) -> bytes:
    doc = Document()

    # 1. Page setup: A4, 18 mm top/bottom, 15 mm left/right
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(18)
    section.bottom_margin = Mm(18)
    section.left_margin = Mm(15)
    section.right_margin = Mm(15)

    # 2. Typography & Complex script setup
    style_info = paper.get('style', {}) or {}
    font_name = style_info.get('fontFamily', 'Times New Roman')
    font_size = style_info.get('fontSize', 11)

    normal_style = doc.styles['Normal']
    normal_style.font.name = font_name
    normal_style.font.size = Pt(font_size)
    normal_style.paragraph_format.line_spacing = 1.35
    normal_style.paragraph_format.space_after = Pt(0)
    normal_style.paragraph_format.space_before = Pt(0)

    rPr = normal_style.element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:cs'), font_name)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    sz_val = str(int(round(font_size * 2)))
    szCs = OxmlElement('w:szCs')
    szCs.set(qn('w:val'), sz_val)
    rPr.append(szCs)

    # 3. Footer setup: Page X of Y
    _setup_footer(section)

    # 4. Header block
    school = paper.get('school', '')
    if school:
        p_school = doc.add_paragraph()
        p_school.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_school.paragraph_format.space_after = Pt(2)
        r_school = p_school.add_run(school)
        r_school.bold = True
        r_school.font.size = Pt(font_size + 3)

    exam = paper.get('exam', '')
    if exam:
        p_exam = doc.add_paragraph()
        p_exam.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_exam.paragraph_format.space_after = Pt(4)
        r_exam = p_exam.add_run(exam)
        r_exam.bold = True
        r_exam.font.size = Pt(font_size + 1)

    # Borderless 3-column table for meta info
    meta_tbl = doc.add_table(rows=1, cols=3)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    _make_borderless(meta_tbl)
    _set_cant_split(meta_tbl.rows[0])

    meta_tbl.columns[0].width = Mm(70)
    meta_tbl.columns[1].width = Mm(55)
    meta_tbl.columns[2].width = Mm(55)

    c_class = meta_tbl.cell(0, 0)
    p_c0 = c_class.paragraphs[0]
    p_c0.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_c0 = p_c0.add_run(f"CLASS - {paper.get('class', '')} {paper.get('subject', '')}")
    r_c0.bold = True

    c_time = meta_tbl.cell(0, 1)
    p_c1 = c_time.paragraphs[0]
    p_c1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_c1 = p_c1.add_run(f"TIME - {paper.get('time', '')}")
    r_c1.bold = True

    c_marks = meta_tbl.cell(0, 2)
    p_c2 = c_marks.paragraphs[0]
    p_c2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_c2 = p_c2.add_run(f"MARKS - {paper.get('totalMarks', '')}")
    r_c2.bold = True

    # 5. Build view model for exact continuous question numbering and marks formula
    view_data = build_view(paper)

    for s in view_data.get('sections', []):
        # Section heading paragraph with right-aligned tab stop and bottom border
        p_sec = doc.add_paragraph()
        p_sec.paragraph_format.space_before = Pt(12)
        p_sec.paragraph_format.space_after = Pt(4)
        p_sec.paragraph_format.keep_with_next = True
        p_sec.paragraph_format.tab_stops.add_tab_stop(Mm(180), WD_TAB_ALIGNMENT.RIGHT)

        pPr = p_sec._p.get_or_add_pPr()
        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="8" w:space="4" w:color="000000"/></w:pBdr>')
        pPr.append(pBdr)

        r_lbl = p_sec.add_run(f"Section {s.get('label', '')}")
        r_lbl.bold = True
        r_lbl.font.size = Pt(font_size + 0.5)

        p_sec.add_run('\t')

        r_mline = p_sec.add_run(s.get('marks_line', ''))
        r_mline.bold = True
        r_mline.font.size = Pt(font_size + 0.5)

        # Optional instruction
        instruction = s.get('instruction', '')
        if instruction:
            p_inst = doc.add_paragraph()
            p_inst.paragraph_format.space_after = Pt(4)
            p_inst.paragraph_format.keep_with_next = True
            r_it = p_inst.add_run(instruction)
            r_it.italic = True

        # Rows (Questions)
        for row in s.get('rows', []):
            no = row['no']
            q = row['q']
            tpl = row['tpl']
            fig = q.get('figure')

            if tpl == 'mcq':
                p_q = doc.add_paragraph()
                p_q.paragraph_format.space_before = Pt(6)
                p_q.paragraph_format.space_after = Pt(2)
                p_q.paragraph_format.keep_with_next = True
                p_q.add_run(f'{no}. ').bold = True
                p_q.add_run(q.get('text', ''))

                if fig:
                    _add_figure(doc, fig, keep_with_next=True)

                valid_opts = [o for o in q.get('options', []) if o and str(o).strip()]
                if valid_opts:
                    letters = 'ABCDEFGH'
                    num_rows = math.ceil(len(valid_opts) / 2)
                    tbl_opts = doc.add_table(rows=num_rows, cols=2)
                    tbl_opts.alignment = WD_TABLE_ALIGNMENT.CENTER
                    _make_borderless(tbl_opts)
                    tbl_opts.columns[0].width = Mm(90)
                    tbl_opts.columns[1].width = Mm(90)

                    for idx, opt in enumerate(valid_opts):
                        r_idx = idx // 2
                        c_idx = idx % 2
                        cell = tbl_opts.cell(r_idx, c_idx)
                        cp = cell.paragraphs[0]
                        cp.paragraph_format.left_indent = Mm(6)
                        cp.paragraph_format.space_after = Pt(2)
                        cp.add_run(f'{letters[idx]}) {opt}')

                    for trow in tbl_opts.rows:
                        _set_cant_split(trow)

            elif tpl == 'sequence':
                p_q = doc.add_paragraph()
                p_q.paragraph_format.space_before = Pt(6)
                p_q.paragraph_format.space_after = Pt(2)
                p_q.paragraph_format.keep_with_next = True
                p_q.add_run(f'{no}. ').bold = True
                p_q.add_run(q.get('text', ''))

                if fig:
                    _add_figure(doc, fig, keep_with_next=True)

                valid_items = [it for it in q.get('items', []) if it and str(it).strip()]
                roman = ['i', 'ii', 'iii', 'iv', 'v', 'vi']
                for idx, it in enumerate(valid_items):
                    p_it = doc.add_paragraph()
                    p_it.paragraph_format.left_indent = Mm(6)
                    p_it.paragraph_format.space_after = Pt(1)
                    p_it.paragraph_format.keep_with_next = True
                    rom = roman[idx] if idx < len(roman) else str(idx + 1)
                    p_it.add_run(f'({rom}) {it}')

                valid_opts = [o for o in q.get('options', []) if o and str(o).strip()]
                if valid_opts:
                    letters = 'ABCDEFGH'
                    num_rows = math.ceil(len(valid_opts) / 2)
                    tbl_opts = doc.add_table(rows=num_rows, cols=2)
                    tbl_opts.alignment = WD_TABLE_ALIGNMENT.CENTER
                    _make_borderless(tbl_opts)
                    tbl_opts.columns[0].width = Mm(90)
                    tbl_opts.columns[1].width = Mm(90)

                    for idx, opt in enumerate(valid_opts):
                        r_idx = idx // 2
                        c_idx = idx % 2
                        cell = tbl_opts.cell(r_idx, c_idx)
                        cp = cell.paragraphs[0]
                        cp.paragraph_format.left_indent = Mm(6)
                        cp.paragraph_format.space_after = Pt(2)
                        cp.add_run(f'{letters[idx]}. {opt}')

                    for trow in tbl_opts.rows:
                        _set_cant_split(trow)

            elif tpl == 'assertion_reason':
                p_q = doc.add_paragraph()
                p_q.paragraph_format.space_before = Pt(6)
                p_q.paragraph_format.space_after = Pt(1)
                p_q.paragraph_format.keep_with_next = True
                p_q.add_run(f'{no}. ').bold = True
                p_q.add_run(f"Assertion (A): {q.get('assertion', '')}")

                p_r = doc.add_paragraph()
                p_r.paragraph_format.space_after = Pt(2)
                p_r.paragraph_format.keep_with_next = True
                p_r.add_run(f"Reason (R): {q.get('reason', '')}")

                # Position rule: directly below Reason line
                if fig:
                    _add_figure(doc, fig, keep_with_next=True)

                valid_opts = [o for o in q.get('options', []) if o and str(o).strip()]
                letters = 'ABCD'
                for idx, opt in enumerate(valid_opts):
                    p_opt = doc.add_paragraph()
                    p_opt.paragraph_format.left_indent = Mm(6)
                    p_opt.paragraph_format.space_after = Pt(1)
                    let = letters[idx] if idx < len(letters) else str(idx + 1)
                    p_opt.add_run(f'{let}) {opt}')

            elif tpl == 'match':
                p_q = doc.add_paragraph()
                p_q.paragraph_format.space_before = Pt(6)
                p_q.paragraph_format.space_after = Pt(2)
                p_q.paragraph_format.keep_with_next = True
                p_q.add_run(f'{no}. ').bold = True
                p_q.add_run(q.get('text') or 'Match the following')

                if fig:
                    _add_figure(doc, fig, keep_with_next=True)

                col_a = q.get('columnA', [])
                col_b = q.get('columnB', [])
                pairs = []
                for i in range(max(len(col_a), len(col_b))):
                    a = col_a[i] if i < len(col_a) else ''
                    b = col_b[i] if i < len(col_b) else ''
                    if (a and str(a).strip()) or (b and str(b).strip()):
                        pairs.append((a, b))

                if pairs:
                    tbl_match = doc.add_table(rows=len(pairs) + 1, cols=2)
                    tbl_match.style = 'Table Grid'
                    tbl_match.alignment = WD_TABLE_ALIGNMENT.CENTER
                    tbl_match.columns[0].width = Mm(80)
                    tbl_match.columns[1].width = Mm(80)

                    # Header row
                    hdr_a = tbl_match.cell(0, 0).paragraphs[0]
                    hdr_a.add_run('Column A').bold = True
                    hdr_b = tbl_match.cell(0, 1).paragraphs[0]
                    hdr_b.add_run('Column B').bold = True

                    letters = 'abcdefgh'
                    for idx, (a, b) in enumerate(pairs):
                        let = letters[idx] if idx < len(letters) else str(idx + 1)
                        tbl_match.cell(idx + 1, 0).paragraphs[0].text = f'{idx + 1}. {a}'
                        tbl_match.cell(idx + 1, 1).paragraphs[0].text = f'({let}) {b}'

                    for trow in tbl_match.rows:
                        _set_cant_split(trow)

                valid_opts = [o for o in q.get('options', []) if o and str(o).strip()]
                if valid_opts:
                    letters = 'ABCD'
                    num_rows = math.ceil(len(valid_opts) / 2)
                    tbl_opts = doc.add_table(rows=num_rows, cols=2)
                    tbl_opts.alignment = WD_TABLE_ALIGNMENT.CENTER
                    _make_borderless(tbl_opts)
                    tbl_opts.columns[0].width = Mm(90)
                    tbl_opts.columns[1].width = Mm(90)

                    for idx, opt in enumerate(valid_opts):
                        r_idx = idx // 2
                        c_idx = idx % 2
                        cell = tbl_opts.cell(r_idx, c_idx)
                        cp = cell.paragraphs[0]
                        cp.paragraph_format.left_indent = Mm(6)
                        cp.paragraph_format.space_after = Pt(2)
                        let = letters[idx] if idx < len(letters) else str(idx + 1)
                        cp.add_run(f'{let}. {opt}')

                    for trow in tbl_opts.rows:
                        _set_cant_split(trow)

            elif tpl == 'text':  # short or long
                p_q = doc.add_paragraph()
                p_q.paragraph_format.space_before = Pt(6)
                p_q.paragraph_format.space_after = Pt(2)
                p_q.paragraph_format.keep_with_next = bool(fig)
                p_q.add_run(f'{no}. ').bold = True
                p_q.add_run(q.get('text', ''))

                if fig:
                    _add_figure(doc, fig, keep_with_next=False)

            elif tpl == 'passage':
                p_q = doc.add_paragraph()
                p_q.paragraph_format.space_before = Pt(6)
                p_q.paragraph_format.space_after = Pt(2)
                p_q.paragraph_format.keep_with_next = True
                p_q.add_run(f'{no}. ').bold = True
                p_q.add_run('Read the following paragraph and answer the questions given below')

                passage_text = q.get('passage', '')
                if passage_text:
                    p_p = doc.add_paragraph()
                    p_p.paragraph_format.space_after = Pt(4)
                    p_p.paragraph_format.keep_with_next = True
                    p_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    p_p.add_run(passage_text)

                # Position rule: directly below passage text
                if fig:
                    _add_figure(doc, fig, keep_with_next=True)

                sub_questions = [sq for sq in q.get('subQuestions', []) if sq.get('text') and str(sq.get('text')).strip()]
                for sq_idx, sq in enumerate(sub_questions):
                    p_sq = doc.add_paragraph()
                    p_sq.paragraph_format.left_indent = Mm(6)
                    p_sq.paragraph_format.tab_stops.add_tab_stop(Mm(180), WD_TAB_ALIGNMENT.RIGHT)
                    p_sq.paragraph_format.space_after = Pt(2)
                    p_sq.add_run(f'{sq_idx + 1}. {sq.get("text", "")}')
                    p_sq.add_run('\t')
                    r_sqm = p_sq.add_run(f'({sq.get("marks", 1)})')
                    r_sqm.bold = True

            elif tpl == 'map':
                p_q = doc.add_paragraph()
                p_q.paragraph_format.space_before = Pt(6)
                p_q.paragraph_format.space_after = Pt(2)
                p_q.paragraph_format.keep_with_next = True
                p_q.add_run(f'{no}. ').bold = True
                p_q.add_run(q.get('text') or 'Locate the following places in an Outline map of India')

                if fig:
                    _add_figure(doc, fig, keep_with_next=True)

                places = [p for p in q.get('places', []) if p and str(p).strip()]
                if places:
                    num_rows = math.ceil(len(places) / 2)
                    tbl_places = doc.add_table(rows=num_rows, cols=2)
                    tbl_places.alignment = WD_TABLE_ALIGNMENT.CENTER
                    _make_borderless(tbl_places)
                    tbl_places.columns[0].width = Mm(90)
                    tbl_places.columns[1].width = Mm(90)

                    for idx, place in enumerate(places):
                        r_idx = idx // 2
                        c_idx = idx % 2
                        cell = tbl_places.cell(r_idx, c_idx)
                        cp = cell.paragraphs[0]
                        cp.paragraph_format.left_indent = Mm(6)
                        cp.paragraph_format.space_after = Pt(2)
                        cp.add_run(f'{idx + 1}. {place}')

                    for trow in tbl_places.rows:
                        _set_cant_split(trow)

            else:
                p_unk = doc.add_paragraph()
                p_unk.paragraph_format.space_before = Pt(4)
                p_unk.paragraph_format.space_after = Pt(2)
                p_unk.add_run(f'{no}. Unsupported question type: {q.get("type", "none")}')
                if fig:
                    _add_figure(doc, fig, keep_with_next=False)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
